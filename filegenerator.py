"""
DOCX Generator สำหรับ Exam Formatter
สร้างไฟล์ Word ตามรูปแบบมาตรฐาน
"""

import os
from docx import Document as DocxDocument
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .config import (
    LOGO_CONFIG, PAGE_CONFIG, GRID_CM, GRID_SPACING,
    FONT_CONFIG, COLOR_CONFIG, DEFAULTS
)
from .image_utils import add_floating_image

def hide_borders(table):
    """ซ่อนเส้นขอบตาราง"""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders = OxmlElement('w:tblBorders')
    
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'none')
        element.set(qn('w:sz'), '0')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), 'auto')
        borders.append(element)
    
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)

def set_cell_margins(cell, top=0, bottom=0, left=100, right=100):
    """ตั้งค่า margin ของ cell"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    
    for edge, val in [('top', top), ('bottom', bottom), ('start', left), ('end', right)]:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    
    tcPr.append(tcMar)

def add_separator_line(doc, color='2E5C8A'):
    """เพิ่มเส้นคั่นแนวนอน"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(0)
    
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    
    return p

def generate_docx(data, output_path, logo_paths):
    """
    สร้างไฟล์ DOCX จากข้อมูลข้อสอบ
    
    Parameters:
    - data: dict ของข้อมูลข้อสอบ
    - output_path: path ที่จะบันทึกไฟล์
    - logo_paths: dict ของ path Logo
    """
    doc = DocxDocument()
    
    # ตั้งค่าหน้ากระดาษ
    section = doc.sections[0]
    section.page_height = Cm(PAGE_CONFIG['height_cm'])
    section.page_width = Cm(PAGE_CONFIG['width_cm'])
    section.top_margin = Cm(PAGE_CONFIG['top_margin_cm'])
    section.bottom_margin = Cm(PAGE_CONFIG['bottom_margin_cm'])
    section.left_margin = Cm(PAGE_CONFIG['left_margin_cm'])
    section.right_margin = Cm(PAGE_CONFIG['right_margin_cm'])
    
    # ========== HEADER ==========
    header = section.header
    
    # เพิ่มเส้นคั่น
    header_sep = header.add_paragraph()
    header_sep.paragraph_format.space_before = Pt(0)
    header_sep.paragraph_format.space_after = Pt(0)
    header_sep.paragraph_format.line_spacing = Pt(0)
    
    pPr = header_sep._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), COLOR_CONFIG['primary'])
    pBdr.append(bottom)
    pPr.append(pBdr)
    
    # เพิ่ม Logo ทั้งหมดใน Header
    for key, config in LOGO_CONFIG.items():
        if config['location'] == 'header':
            if logo_paths.get(key) and os.path.exists(logo_paths[key]):
                add_floating_image(
                    container=header,
                    image_path=logo_paths[key],
                    width_cm=config['width_cm'],
                    height_cm=config['height_cm'],
                    h_pos_cm=config['h_pos_cm'],
                    v_pos_cm=config['v_pos_cm'],
                    behind_doc=config.get('behind_doc', True)
                )

    # ========== TOP SECTION ==========
    top_tbl = doc.add_table(2, 2)
    hide_borders(top_tbl)
    top_tbl.rows[0].cells[0].width = Cm(8)
    top_tbl.rows[0].cells[1].width = Cm(8)
    
    for row in top_tbl.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=60, bottom=60, left=150, right=150)
            cell.paragraphs[0].text = ""

    p_space_top = doc.add_paragraph()
    p_space_top.paragraph_format.space_before = Pt(0)
    p_space_top.paragraph_format.space_after = Cm(0.64)
    p_space_top.paragraph_format.line_spacing = Pt(0)

    # ========== INFO SECTION ==========
    info_items = [
        ("หน่วยการเรียนรู้", data.get('unit', DEFAULTS['unit'])),
        ("เรื่อง", data.get('topic', DEFAULTS['topic'])),
        ("ระดับชั้น", data.get('grade', DEFAULTS['grade'])),
        ("ชื่อ", ".................................................................................  ชั้น: ............  เลขที่: ............")
    ]
    
    for label, value in info_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Cm(GRID_CM)
        p.paragraph_format.line_spacing = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4))
        
        run_label = p.add_run(label)
        run_label.bold = True
        run_label.font.size = Pt(FONT_CONFIG['size']['info_label'])
        run_label.font.name = FONT_CONFIG['main']
        run_label.font.color.rgb = COLOR_CONFIG['primary_rgb']
        
        p.add_run("\t")
        
        run_value = p.add_run(value)
        run_value.font.size = Pt(FONT_CONFIG['size']['info_value'])
        run_value.font.name = FONT_CONFIG['main']

    # ========== คำชี้แจง ==========
    p_instr = doc.add_paragraph()
    p_instr.paragraph_format.space_before = Pt(0)
    p_instr.paragraph_format.space_after = Cm(GRID_CM * 2)
    p_instr.paragraph_format.line_spacing = Pt(0)
    p_instr.paragraph_format.tab_stops.add_tab_stop(Cm(4))
    
    run_label = p_instr.add_run("คำชี้แจง")
    run_label.bold = True
    run_label.font.size = Pt(FONT_CONFIG['size']['info_label'])
    run_label.font.name = FONT_CONFIG['main']
    run_label.font.color.rgb = COLOR_CONFIG['primary_rgb']
    
    p_instr.add_run("\t")
    
    run_value = p_instr.add_run(data.get('instruction', DEFAULTS['instruction']))
    run_value.font.size = Pt(FONT_CONFIG['size']['info_value'])
    run_value.font.name = FONT_CONFIG['main']

    # ========== เส้นคั่น ==========
    sep_line = add_separator_line(doc, color=COLOR_CONFIG['primary'])
    sep_line.paragraph_format.space_before = Pt(0)
    sep_line.paragraph_format.space_after = Cm(GRID_CM * 2)

    # ========== ข้อสอบ ==========
    for q in data.get('questions', []):
        # โจทย์
        p_q = doc.add_paragraph()
        p_q.paragraph_format.line_spacing = 1.0
        p_q.paragraph_format.space_before = Pt(0)
        p_q.paragraph_format.space_after = Cm(GRID_SPACING['between_question_and_choice'] * GRID_CM)
        
        tab_stop_pos = GRID_SPACING['question_number_gap'] * GRID_CM
        p_q.paragraph_format.tab_stops.add_tab_stop(Cm(tab_stop_pos), alignment=WD_ALIGN_PARAGRAPH.LEFT)
        
        run_no = p_q.add_run(f"{q['no']}.")
        run_no.bold = True
        run_no.font.size = Pt(FONT_CONFIG['size']['question'])
        run_no.font.name = FONT_CONFIG['main']
        
        p_q.add_run("\t")
        
        run_text = p_q.add_run(q['text'])
        run_text.font.size = Pt(FONT_CONFIG['size']['question'])
        run_text.font.name = FONT_CONFIG['main']
        
        # ตัวเลือก
        available_letters = [l for l in ['ก', 'ข', 'ค', 'ง'] if l in q['choices']]
        last_letter = available_letters[-1] if available_letters else None
        
        for letter in available_letters:
            p_c = doc.add_paragraph()
            p_c.paragraph_format.line_spacing = 1.0
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.left_indent = Cm(GRID_SPACING['choice_indent'] * GRID_CM)
            
            if letter == last_letter:
                p_c.paragraph_format.space_after = Cm(GRID_SPACING['after_last_choice'] * GRID_CM)
            else:
                p_c.paragraph_format.space_after = Cm(GRID_SPACING['between_choices'] * GRID_CM)
            
            run_choice = p_c.add_run(f"{letter}.  {q['choices'][letter]}")
            run_choice.font.size = Pt(FONT_CONFIG['size']['choice'])
            run_choice.font.name = FONT_CONFIG['main']
            
            if letter == 'ก':
                run_choice.bold = True
    
    # ========== FOOTER ==========
    footer = section.footer
    
    # เพิ่ม Logo ใน Footer
    for key, config in LOGO_CONFIG.items():
        if config['location'] == 'footer':
            if logo_paths.get(key) and os.path.exists(logo_paths[key]):
                add_floating_image(
                    container=footer,
                    image_path=logo_paths[key],
                    width_cm=config['width_cm'],
                    height_cm=config['height_cm'],
                    h_pos_cm=config['h_pos_cm'],
                    v_pos_cm=config['v_pos_cm'],
                    behind_doc=config.get('behind_doc', True)
                )
    
    # บันทึกไฟล์
    doc.save(output_path)
